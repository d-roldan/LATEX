"""Exporta los cuatro manuales Markdown a Word con sus capturas locales.

Uso: python scripts/exportar_manuales_word.py
Requiere python-docx y Pillow.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "word"
PROFILES = ("FABRICACION", "LABORATORIO", "ENVASADO", "ADMINISTRADOR")
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
IMAGE = re.compile(r"^!\[([^]]*)\]\(([^)]+)\)$")
ITEM = re.compile(r"^(\s*)(\d+\.|-)\s+(.*)$")


def set_font(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(5)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    set_font(styles["Normal"], 10.5)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    set_font(styles["Title"], 18, True)
    styles["Title"].paragraph_format.space_after = Pt(9)
    title_properties = styles["Title"]._element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    for name, size, before, after in (
        ("Heading 1", 14, 13, 5),
        ("Heading 2", 12, 10, 4),
    ):
        style = styles[name]
        set_font(style, size, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = styles["Normal"]
    footer.add_run("Página ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)


def inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def add_image(doc, source_dir, alt, relative):
    path = (source_dir / relative).resolve()
    if not path.is_file() or not path.is_relative_to(DOCS.resolve()):
        raise FileNotFoundError(f"Imagen no disponible: {path}")
    with Image.open(path) as picture:
        px_width, px_height = picture.size
    width = min(6.85, 4.45 * px_width / px_height)
    if px_width < 400:
        width = min(width, max(1.3, px_width / 130))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = bool(alt)
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    # Texto alternativo accesible en Word.
    shape._inline.docPr.set("descr", alt)
    if alt:
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(8)
        caption.paragraph_format.keep_together = True
        caption_run = caption.add_run(alt)
        caption_run.font.size = Pt(8)
        caption_run.font.color.rgb = RGBColor(70, 70, 70)


def add_table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = True
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            inline(cell.paragraphs[0], value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
                if row_index == 0:
                    run.bold = True
        if row_index == 0:
            table.rows[row_index]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    table._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_text(doc, lines):
    text = " ".join(line.strip() for line in lines)
    paragraph = doc.add_paragraph()
    inline(paragraph, text)


def convert(source):
    doc = Document()
    configure(doc)
    lines = source.read_text(encoding="utf-8").splitlines()
    pending = []
    index = 0

    def flush():
        if pending:
            add_text(doc, pending)
            pending.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush()
            index += 1
            continue
        if stripped == "---":
            flush()
            index += 1
            continue
        image = IMAGE.match(stripped)
        if image:
            flush()
            add_image(doc, source.parent, image.group(1), image.group(2))
            index += 1
            continue
        if stripped.startswith("|"):
            flush()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue
        if stripped.startswith("### ") or stripped.startswith("## ") or stripped.startswith("# "):
            flush()
            if stripped.startswith("# "):
                title = stripped[2:].replace(" · ", " de ")
                doc.add_paragraph(title, "Title")
            else:
                level = 2 if stripped.startswith("### ") else 1
                doc.add_paragraph(stripped[level + 2 :], f"Heading {level}")
            index += 1
            continue
        item = ITEM.match(line)
        if item:
            flush()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2 + len(item.group(1)) * 0.13)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(2.5)
            prefix = "• " if item.group(2) == "-" else f"{item.group(2)} "
            paragraph.add_run(prefix)
            inline(paragraph, item.group(3))
            index += 1
            continue
        if stripped.startswith("> "):
            flush()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            inline(paragraph, stripped[2:])
            index += 1
            continue
        pending.append(line)
        index += 1
    flush()
    doc.core_properties.title = source.stem.replace("MANUAL_USUARIO_", "Manual de usuario de ").title()
    doc.core_properties.subject = "Manual de usuario de DISAL Planta de Látex"
    output = OUTPUT / f"{source.stem}.docx"
    doc.save(output)
    return output


def main():
    OUTPUT.mkdir(exist_ok=True)
    for profile in PROFILES:
        source = DOCS / f"MANUAL_USUARIO_{profile}.md"
        output = convert(source)
        print(output)


if __name__ == "__main__":
    main()
